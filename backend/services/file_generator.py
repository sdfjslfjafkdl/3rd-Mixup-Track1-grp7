"""
파일 생성기: 제안서(docx)와 평가 매트릭스(xlsx)를 생성
"""
import re
from docx import Document
from docx.enum.text import WD_PARAGRAPH_ALIGNMENT
from openpyxl import Workbook
import os
from datetime import datetime

OUTPUT_DIR = os.path.abspath(os.path.join(os.path.dirname(__file__), '..', '..', 'output'))
if not os.path.exists(OUTPUT_DIR):
    os.makedirs(OUTPUT_DIR, exist_ok=True)


def _plain_text(value: str) -> str:
    text = value.replace("**", "").replace("__", "").strip()
    text = re.sub(r"`([^`]+)`", r"\1", text)
    return text


def generate_docx(proposal: dict, candidate: dict, profile: dict) -> str:
    """제안서 초안 텍스트를 docx로 저장하고 경로 반환"""
    doc = Document()
    title = candidate.get("title") or f"Proposal - {candidate.get('id')}"
    heading = doc.add_heading(title, level=1)
    heading.alignment = WD_PARAGRAPH_ALIGNMENT.CENTER

    agency = candidate.get("agency") or "발주기관 정보 없음"
    proposal_company = profile.get("name") or "제안사 정보 없음"
    header = doc.add_paragraph()
    header.add_run(f"제안사: {proposal_company}\n").bold = True
    header.add_run(f"발주기관: {agency}\n")
    header.add_run(f"생성일: {datetime.now().strftime('%Y-%m-%d %H:%M:%S')}")
    doc.add_paragraph("")

    # draft text (멀티라인)
    draft = proposal.get("draft_text", "")
    for raw_line in draft.splitlines():
        line = raw_line.strip()
        if not line:
            doc.add_paragraph("")
            continue
        if line.startswith("# "):
            doc.add_heading(_plain_text(line[2:].strip()), level=1)
            continue
        if line.startswith("## "):
            doc.add_heading(_plain_text(line[3:].strip()), level=2)
            continue
        if line.startswith("### "):
            doc.add_heading(_plain_text(line[4:].strip()), level=3)
            continue
        if line.startswith("- "):
            doc.add_paragraph(_plain_text(line[2:].strip()), style="List Bullet")
            continue
        doc.add_paragraph(_plain_text(line))

    # 같은 공고를 여러 번 revise 할 때 초 단위 timestamp로는 충돌할 수 있어 ms 단위 사용
    filename = f"proposal_{candidate.get('id')}_{int(datetime.now().timestamp() * 1000)}.docx"
    path = os.path.join(OUTPUT_DIR, filename)
    doc.save(path)
    return path


def generate_xlsx(proposal: dict, match_result: dict, assessment: dict, candidate: dict) -> str:
    """평가 매트릭스 및 스코어를 xlsx로 저장하고 경로 반환"""
    wb = Workbook()
    ws = wb.active
    ws.title = "Scorecard"

    # 헤더
    ws.append(["요구사항 ID", "요구사항", "가중치", "응답", "항목 점수"])

    matrix = proposal.get('matrix', [])
    details = match_result.get('details', [])
    # 매트릭스 및 점수 기록
    for row in matrix:
        req_id = row.get('requirement_id')
        req_text = row.get('requirement_text')
        weight = row.get('weight')
        response = row.get('response')
        # find score from details
        score = next((d.get('score') for d in details if d.get('req_id') == req_id), None)
        ws.append([req_id, req_text, weight, response, score])

    # 요약 시트
    ws2 = wb.create_sheet(title="Summary")
    ws2.append(["공고 ID", candidate.get('id')])
    ws2.append(["공고명", candidate.get('title')])
    ws2.append(["전체 적합도", match_result.get('overall_score')])
    ws2.append(["Win Probability", f"{match_result.get('win_probability')}%"])
    ws2.append(["자격 매칭 점수", assessment.get('match_score')])
    ws2.append(["자격 충족 여부", assessment.get('eligible')])

    # 같은 공고를 여러 번 revise 할 때 초 단위 timestamp로는 충돌할 수 있어 ms 단위 사용
    filename = f"evaluation_{candidate.get('id')}_{int(datetime.now().timestamp() * 1000)}.xlsx"
    path = os.path.join(OUTPUT_DIR, filename)
    wb.save(path)
    return path
